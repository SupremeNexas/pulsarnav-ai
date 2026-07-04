from __future__ import annotations

import re
import subprocess
import os
import ast
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS
TODAY = date.today().isoformat()

BLUE = RGBColor(31, 78, 121)
ACCENT = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 89, 89)
LIGHT = "F2F4F7"
BORDER = "D9E2EC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)


def add_header_footer(doc: Document, label: str) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = label
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(header.runs[0], size=8.5, color=MUTED)
        footer = section.footer.paragraphs[0]
        footer.text = f"PulsarNav AI | {TODAY}"
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(footer.runs[0], size=8.5, color=MUTED)


def add_cover(doc: Document, title: str, subtitle: str, kind: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(kind.upper())
    set_run_font(run, size=10, color=MUTED, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(title)
    set_run_font(run, size=25, color=RGBColor(0, 0, 0), bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(subtitle)
    set_run_font(run, size=13, color=MUTED)

    metadata = [
        ("Project", "PulsarNav AI"),
        ("Repository", "/Users/supryo/Desktop/pulsar"),
        ("Prepared", TODAY),
        ("Document Class", kind),
        ("Primary Audience", "Academic reviewers, aerospace mentors, flight dynamics developers"),
    ]
    add_key_value_table(doc, metadata, widths=(1.45, 5.25))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(
        "This document describes PulsarNav AI as a scientific software system: "
        "why each component exists, the mathematics behind it, how it is implemented, "
        "and how its outcomes are validated."
    )
    set_run_font(run, size=10.5, color=RGBColor(40, 40, 40))
    doc.add_page_break()


def add_key_value_table(doc: Document, rows: list[tuple[str, str]], widths=(1.7, 4.8)) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    for row_idx, (key, value) in enumerate(rows):
        cells = table.rows[row_idx].cells
        for i, width in enumerate(widths):
            cells[i].width = Inches(width)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cells[0], LIGHT)
        cells[0].paragraphs[0].add_run(key).bold = True
        cells[1].paragraphs[0].add_run(value)


def add_normal_paragraph(doc: Document, text: str) -> None:
    text = cleanup_inline(text)
    if not text:
        return
    p = doc.add_paragraph()
    for chunk, is_code in split_inline_code(text):
        run = p.add_run(chunk)
        if is_code:
            set_run_font(run, size=9.3, name="Consolas", color=BLUE)


def split_inline_code(text: str):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            yield part[1:-1], True
        elif part:
            yield part, False


def cleanup_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    replacements = {
        r"\mathbf": "",
        r"\vec": "",
        r"\hat": "",
        r"\text": "",
        r"\left": "",
        r"\right": "",
        r"\cdot": " dot ",
        r"\times": " x ",
        r"\frac": "frac",
        r"\sqrt": "sqrt",
        r"\sum": "sum",
        r"\int": "integral",
        r"\Delta": "Delta",
        r"\delta": "delta",
        r"\lambda": "lambda",
        r"\phi": "phi",
        r"\sigma": "sigma",
        r"\mu": "mu",
        r"\eta": "eta",
        r"\tau": "tau",
        r"\approx": "~",
        r"\ge": ">=",
        r"\le": "<=",
        r"\sim": "~",
        r"\bmod": "mod",
        r"\begin": "begin",
        r"\end": "end",
        r"\vdots": "...",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace("{", "").replace("}", "")
    text = text.replace("<br>", " ")
    return text.strip()


def add_bullet(doc: Document, text: str, numbered=False) -> None:
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(cleanup_inline(text))


def parse_markdown_table(block: list[str]) -> list[list[str]] | None:
    rows = []
    for line in block:
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            return None
        cells = [cleanup_inline(c.strip()) for c in stripped.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
            continue
        rows.append(cells)
    return rows if len(rows) >= 2 else None


def add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx].cells
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=8.7 if len(text) > 80 else 9.2, bold=(r_idx == 0))
            set_cell_margins(cells[c_idx], top=100, bottom=100, start=120, end=120)
            cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cells[c_idx], LIGHT)
        if r_idx == 0:
            set_repeat_table_header(table.rows[r_idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(style="Code Block")
        p.add_run(cleanup_inline(line.rstrip()))


def add_equation_block(doc: Document, lines: list[str]) -> None:
    cleaned = [cleanup_inline(line).strip("$").strip() for line in lines if cleanup_inline(line).strip("$").strip()]
    if not cleaned:
        return
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.15)
    p.add_run(" ".join(cleaned))


def render_markdown(doc: Document, markdown: str, max_heading_level=3) -> None:
    lines = markdown.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_lines[:90])
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("$$"):
            equation_lines = []
            current = line.strip()
            if current != "$$":
                equation_lines.append(current.strip("$"))
                i += 1
            else:
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    equation_lines.append(lines[i])
                    i += 1
                if i < len(lines):
                    closing = lines[i].strip()
                    if closing != "$$":
                        equation_lines.append(closing.strip("$"))
                    i += 1
            add_equation_block(doc, equation_lines)
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            table_rows = parse_markdown_table(block)
            if table_rows:
                add_table(doc, table_rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), max_heading_level)
            text = cleanup_inline(heading.group(2))
            doc.add_heading(text, level=level)
            i += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet:
            add_bullet(doc, bullet.group(1), numbered=False)
            i += 1
            continue
        if numbered:
            add_bullet(doc, numbered.group(1), numbered=True)
            i += 1
            continue
        add_normal_paragraph(doc, line)
        i += 1


def add_static_toc(doc: Document, items: list[str]) -> None:
    doc.add_heading("Document Map", level=1)
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
    doc.add_page_break()


def main_document() -> Document:
    doc = Document()
    style_document(doc)
    add_cover(
        doc,
        "PulsarNav AI Technical Design Document",
        "Autonomous Pulsar-Based Deep Space Navigation Framework",
        "Technical Design Document",
    )
    add_header_footer(doc, "PulsarNav AI Technical Design Document")
    add_static_toc(
        doc,
        [
            "Executive documentation and project outcomes",
            "Mathematical foundations",
            "Software architecture and module descriptions",
            "Data pipeline and validation approach",
            "User interface and visualization documentation",
            "Future work, references, and appendices",
        ],
    )

    pipeline = (ROOT / "PROJECT_PIPELINE.md").read_text()
    render_markdown(doc, pipeline)

    doc.add_page_break()
    doc.add_heading("17. Deliverable-Oriented Achievement Matrix", level=1)
    add_table(
        doc,
        [
            ["Achievement", "Implemented Evidence", "Validation / Output"],
            ["GPS-independent spacecraft navigation framework", "Python modules and Next.js Navigation Lab solve spacecraft position from pulsar timing delays.", "Monte Carlo trials, dashboard telemetry, CSV exports."],
            ["NANOGrav 12.5-year dataset integration", "TEMPO/TEMPO2 PAR/TIM parsing pipeline and catalog/statistics generation.", "Catalog, unit-vector, TOA database, statistics, and ranked pulsar CSVs."],
            ["NHPP signal modeling", "RateFunction and photon TOA generation use non-homogeneous Poisson process logic.", "Delay-estimator tests and CRLB comparison surfaces."],
            ["Epoch folding", "Photon TOAs are folded into phase histograms with velocity-error tolerance checks.", "Profile recovery, phase histograms, and estimator inputs."],
            ["Delay estimation", "CC, NLS, and MLE phase estimators are implemented for pulse-delay recovery.", "Pulse delay tests and compatibility wrapper."],
            ["Navigation geometry", "Direction matrices, measurement matrices, GDOP, and observability utilities.", "Pulsar set optimization and rank/condition diagnostics."],
            ["LS/WLS/EKF estimation", "Linear position solvers, 10-state error-state KF, and 8-state absolute EKF.", "Navigation Lab API, trajectory visualization, and unit tests."],
            ["3D visualization", "React Three Fiber and Three.js spacecraft, Earth/Moon, pulsar, line-of-sight, and trajectory rendering.", "Browser-tested trajectory playback and label collision handling."],
            ["Comparative navigation lab", "DSN, XNAV, and Hybrid EKF simulator in TypeScript.", "Generated comparison report and mission KPI tables."],
        ],
    )

    doc.add_heading("18. Submission Roadmap", level=1)
    for item in [
        "Phase 1: Freeze this design document as the technical baseline for project submission.",
        "Phase 2: Insert final screenshots, generated plots, and Monte Carlo result tables from the latest run.",
        "Phase 3: Add appendix-level source listings only where required by academic submission rules.",
        "Phase 4: Convert the executive summary into an oral-defense abstract and presentation script.",
    ]:
        add_bullet(doc, item, numbered=True)

    doc.add_page_break()
    doc.add_heading("Appendix A. Equation Implementation Map", level=1)
    render_markdown(doc, (DOCS / "book_equation_map.md").read_text(), max_heading_level=3)

    doc.add_page_break()
    doc.add_heading("Appendix B. Navigation Engine Reference", level=1)
    render_markdown(doc, (DOCS / "NAVIGATION_ENGINE.md").read_text(), max_heading_level=3)

    doc.add_page_break()
    doc.add_heading("Appendix C. Source Directory Structure", level=1)
    render_markdown(doc, (ROOT / "README.md").read_text().split("## Installation")[0], max_heading_level=3)

    doc.add_page_break()
    doc.add_heading("Appendix D. Python Module Inventory", level=1)
    add_normal_paragraph(
        doc,
        "This appendix lists the primary public functions and classes discovered in the "
        "`pulsar_nav` package. It is intended as a quick engineering reference for reviewers "
        "and future contributors."
    )
    add_module_inventory(doc)

    return doc


def add_module_inventory(doc: Document) -> None:
    rows = [["Module", "Classes", "Functions"]]
    for path in sorted((ROOT / "pulsar_nav").glob("*.py")):
        tree = ast.parse(path.read_text())
        classes = []
        funcs = []
        for node in tree.body:
            if isinstance(node, ast.ClassDef) and not node.name.startswith("_"):
                classes.append(node.name)
            if isinstance(node, ast.FunctionDef) and not node.name.startswith("_"):
                funcs.append(node.name)
        rows.append(
            [
                path.name,
                ", ".join(classes[:8]) or "-",
                ", ".join(funcs[:12]) or "-",
            ]
        )
    add_table(doc, rows)


EXECUTIVE_MD = """
# Executive Summary

PulsarNav AI is a research-grade software system for autonomous spacecraft navigation using millisecond pulsars as natural celestial beacons. The system demonstrates how a spacecraft can estimate its position without GPS and without continuous ground-station tracking by measuring pulsar time-of-arrival delays and solving the resulting navigation geometry.

## Problem and Motivation

GPS does not provide reliable deep-space coverage because GPS satellites are Earth-orbiting transmitters with antenna patterns and signal strengths optimized for terrestrial and near-Earth users. Deep-space missions instead depend primarily on the Deep Space Network, which is accurate but scheduled, ground-dependent, and affected by communication latency. Pulsar navigation offers a complementary autonomous technique: the spacecraft observes stable astronomical sources distributed across the sky and uses timing residuals as navigation measurements.

## What Was Built

- A Python scientific package for catalog parsing, coordinate conversion, timing correction, signal modeling, epoch folding, delay estimation, CRLB analysis, least-squares estimation, and Kalman filtering.
- A Next.js mission-control dashboard for interactive simulation, visualization, export, and comparison.
- A Navigation Lab API that runs LS, WLS, and EKF-style navigation studies.
- A comparative aerospace navigation engine that evaluates DSN-only, XNAV-only, and Hybrid EKF strategies under common flight dynamics.
- 3D spacecraft trajectory playback with true trajectory, estimated trajectory, error vector, pulsar directions, and mission telemetry.

## Core Mathematics

The central observable is the geometric timing delay:

Delta t_i = (r dot n_i) / c

where r is the spacecraft position vector, n_i is the unit direction to pulsar i, and c is the speed of light. With four or more pulsars, the system estimates three position components plus clock bias. Least squares, weighted least squares, and Kalman filtering extend this basic relation to noisy, overdetermined, and time-evolving mission cases.

## Scientific Outcomes

PulsarNav AI integrates real pulsar timing file formats, models photon arrivals as a non-homogeneous Poisson process, reconstructs pulse profiles through epoch folding, estimates phase delays using CC/NLS/MLE methods, and validates navigation performance with Monte Carlo simulation. The web platform turns these algorithms into a usable research console with charts, exports, and 3D flight-state visualization.

## Recommended Use

The project is best presented as a scientific software system rather than a simple code project. Its value is the full chain from astrophysical data to mathematical estimator to validated aerospace visualization.
"""


PRESENTATION_MD = """
# Presentation-Ready Summary

## 1. Title

PulsarNav AI: Autonomous Pulsar-Based Deep Space Navigation Framework.

## 2. Problem

Deep-space spacecraft cannot depend on GPS. Ground-based DSN tracking is powerful but scheduled, delayed, and operationally expensive. Autonomous navigation is required for future lunar, Martian, and interplanetary missions.

## 3. Core Idea

Millisecond pulsars behave like highly stable natural clocks. A spacecraft measures pulse arrival delays from multiple pulsars. Those delays become geometric constraints on spacecraft position.

## 4. Key Equation

Delta t_i = (r dot n_i) / c

This equation maps position into timing delay. The inverse problem recovers spacecraft position and clock bias from multiple pulsar measurements.

## 5. Data Pipeline

NANOGrav PAR/TIM files are parsed into a pulsar catalog, TOA database, timing statistics, ranked pulsar lists, unit vectors, simulation outputs, plots, and dashboard-ready exports.

## 6. Signal Processing

The project models photon arrivals with a non-homogeneous Poisson process, folds photons into phase histograms, and estimates delays using cross-correlation, nonlinear least squares, and maximum likelihood estimation.

## 7. Navigation Solvers

The system includes least squares, weighted least squares, a 10-state error-state Kalman filter, and an 8-state absolute EKF with RK4 propagation and J2 perturbation support.

## 8. Validation

Monte Carlo simulation measures position error across timing noise levels, pulsar counts, and mission regions. CRLB analysis provides theoretical estimator limits.

## 9. Dashboard

The Next.js dashboard provides simulation controls, charts, CSV/PDF style reporting, 3D space visualization, and animated trajectory playback.

## 10. Comparative Lab

The comparison engine evaluates DSN-only, XNAV-only, and Hybrid EKF navigation with DSN range/Doppler, pulsar measurements, and fused estimates.

## 11. Achievements

- GPS-independent spacecraft navigation simulation.
- Real pulsar catalog and timing-data integration.
- Mathematical implementation of pulsar signal, delay, and navigation geometry.
- Research dashboard with interactive aerospace visualization.
- Validation tools suitable for project review and future publication work.

## 12. Limitations

Some UI panels remain experimental, full ATNF/JPL ephemeris integration is future work, and real detector physics can be made more detailed.

## 13. Future Work

Add ATNF catalog ingestion, JPL DE440 ephemerides, covariance ellipsoids, real detector response models, star-tracker fusion, optical navigation, and production-grade hybrid navigation studies.

## 14. Closing

PulsarNav AI demonstrates the full research chain from astrophysics to flight-dynamics software: pulsar timing data, signal processing, estimator mathematics, Monte Carlo validation, and mission-control visualization.
"""


def short_document(title: str, subtitle: str, kind: str, markdown: str) -> Document:
    doc = Document()
    style_document(doc)
    add_cover(doc, title, subtitle, kind)
    add_header_footer(doc, title)
    render_markdown(doc, markdown)
    return doc


def save_doc(doc: Document, filename: str) -> Path:
    path = OUT / filename
    doc.save(path)
    return path


def convert_to_pdf(docx_path: Path) -> Path:
    soffice = "/Users/supryo/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/soffice"
    profile_dir = Path("/private/tmp/pulsarnav-lo-profile")
    profile_dir.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["HOME"] = "/private/tmp"
    subprocess.run(
        [
            soffice,
            "--headless",
            f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        cwd=ROOT,
        env=env,
    )
    return docx_path.with_suffix(".pdf")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    outputs = [
        save_doc(main_document(), "PulsarNav_AI_Technical_Design_Document.docx"),
        save_doc(
            short_document(
                "PulsarNav AI Executive Summary",
                "2-3 page overview of objectives, achievements, architecture, and outcomes",
                "Executive Summary",
                EXECUTIVE_MD,
            ),
            "PulsarNav_AI_Executive_Summary.docx",
        ),
        save_doc(
            short_document(
                "PulsarNav AI Presentation Summary",
                "Slide-adaptable project narrative and technical talking points",
                "Presentation-Ready Summary",
                PRESENTATION_MD,
            ),
            "PulsarNav_AI_Presentation_Summary.docx",
        ),
    ]
    for path in outputs:
        convert_to_pdf(path)
        print(path)


if __name__ == "__main__":
    main()
